from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from backend.app.rag.ingestion.text_cleaner import clean_page_text


def parse_document_bytes(
    *,
    content: bytes,
    filename: str,
    content_type: str,
) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if content_type == "application/pdf" or extension == ".pdf":
        return parse_pdf_bytes(content)
    if (
        content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or extension == ".docx"
    ):
        return parse_docx_bytes(content)
    return parse_text_bytes(content)


def parse_pdf_bytes(content: bytes) -> list[dict]:
    for parser in (_parse_pdf_with_pymupdf, _parse_pdf_with_pdfplumber, _parse_pdf_with_pypdf):
        pages = parser(content)
        if pages:
            return pages
    return []


def _parse_pdf_with_pymupdf(content: bytes) -> list[dict]:
    try:
        import fitz
    except ImportError:
        return []
    pages: list[dict] = []
    with fitz.open(stream=content, filetype="pdf") as document:
        for index, page in enumerate(document, start=1):
            text = clean_page_text(page.get_text("text") or "")
            if text:
                pages.append({"page_number": index, "text": text})
    return pages


def _parse_pdf_with_pdfplumber(content: bytes) -> list[dict]:
    try:
        import pdfplumber
    except ImportError:
        return []
    pages: list[dict] = []
    with pdfplumber.open(BytesIO(content)) as document:
        for index, page in enumerate(document.pages, start=1):
            text = clean_page_text(page.extract_text() or "")
            if text:
                pages.append({"page_number": index, "text": text})
    return pages


def _parse_pdf_with_pypdf(content: bytes) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    pages: list[dict] = []
    for index, page in enumerate(reader.pages, start=1):
        text = clean_page_text(page.extract_text() or "")
        if text:
            pages.append({"page_number": index, "text": text})
    return pages


def parse_text_bytes(content: bytes) -> list[dict]:
    text = content.decode("utf-8", errors="ignore")
    text = clean_page_text(text)
    return [{"page_number": 1, "text": text}] if text else []


def parse_docx_bytes(content: bytes) -> list[dict]:
    try:
        from docx import Document
    except ImportError as exc:
        return _parse_docx_zip_bytes(content)

    document = Document(BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = clean_page_text(paragraph.text or "")
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [clean_page_text(cell.text or "") for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                blocks.append(line)
    text = clean_page_text("\n".join(blocks))
    return [{"page_number": 1, "text": text}] if text else []


def _parse_docx_zip_bytes(content: bytes) -> list[dict]:
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(BytesIO(content)) as archive:
        xml_content = archive.read("word/document.xml")
    root = ET.fromstring(xml_content)
    paragraphs = []
    for paragraph in root.findall(".//w:p", namespace):
        runs = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
        text = clean_page_text("".join(runs))
        if text:
            paragraphs.append(text)
    text = clean_page_text("\n".join(paragraphs))
    return [{"page_number": 1, "text": text}] if text else []
